from fastapi import FastAPI, Depends, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy.orm import Session
import hashlib
import urllib.parse
import io
import calendar
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import models
from database import engine, get_db

app = FastAPI(title="منظومة إدارة المحافظ والمشاريع الاستثمارية")
app.add_middleware(SessionMiddleware, secret_key="secure_investment_system_2026")

# إنشاء الجداول الأساسية
models.Base.metadata.create_all(bind=engine)


# ترقية تلقائية لقاعدة البيانات لإضافة الأعمدة الجديدة دون مسح البيانات السابقة
def auto_migrate_sqlite():
    conn = engine.raw_connection()
    cur = conn.cursor()
    columns_to_ensure = [
        ("users", "role", "TEXT DEFAULT 'admin'"),
        ("users", "investor_id", "INTEGER"),
        ("project_units", "cost_price", "FLOAT DEFAULT 0.0"),
        ("project_units", "target_price", "FLOAT DEFAULT 0.0"),
        ("project_units", "area", "FLOAT DEFAULT 0.0"),
        ("project_units", "unit_type", "TEXT DEFAULT 'شقة'"),
        ("project_units", "status", "TEXT DEFAULT 'متاحة'"),
        ("project_units", "notes", "TEXT"),
        ("sales", "unit_id", "INTEGER"),
        ("sales", "cost_price", "FLOAT DEFAULT 0.0"),
        ("sales", "net_unit_profit", "FLOAT DEFAULT 0.0"),
        ("sales", "sale_type", "TEXT DEFAULT 'تقسيط'"),
        ("sales", "down_payment", "FLOAT DEFAULT 0.0"),
        ("sales", "remaining_amount", "FLOAT DEFAULT 0.0"),
        ("sales", "installments_count", "INTEGER DEFAULT 0"),
        ("sales", "period_type", "TEXT DEFAULT 'شهري'"),
        ("contributions", "mgmt_fee_pct", "FLOAT DEFAULT 0.0"),
    ]
    for tbl, col, cdef in columns_to_ensure:
        try:
            cur.execute(f"PRAGMA table_info({tbl})")
            existing = {r[1] for r in cur.fetchall()}
            if col not in existing:
                cur.execute(f"ALTER TABLE {tbl} ADD COLUMN {col} {cdef}")
                conn.commit()
        except Exception:
            pass
    conn.close()


auto_migrate_sqlite()

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")


def is_auth(request: Request) -> bool:
    return request.session.get("user") is not None


def get_role(request: Request) -> str:
    return request.session.get("role", "admin")


def add_months(sourcedate, months):
    month = sourcedate.month - 1 + months
    year = sourcedate.year + month // 12
    month = month % 12 + 1
    day = min(sourcedate.day, calendar.monthrange(year, month)[1])
    return datetime(year, month, day)


PROFIT_TYPES = ["أرباح معتمدة", "إقرار أرباح", "توزيع دفعة من المطور", "أرباح بيع وحدة"]
CASH_OUT_TYPES = ["مسحوب للمساهم", "مسحوبات نقدية"]
ADMIN_CASH_TYPES = ["مسحوب للإدارة"]


def calc_project_financials(project: models.Project, db: Session):
    if not project:
        return {
            "capital": 0.0, "actual_col": 0.0, "col_from_developer": 0.0, "col_from_sales": 0.0,
            "total_exp": 0.0, "net_profit": 0.0, "total_profits_allocated": 0.0,
            "cash_withdrawn_investors": 0.0, "recovery_pct": 0.0, "remaining_to_recover": 0.0,
            "incoming_transfers": 0.0, "outgoing_transfers": 0.0, "remaining_liquidity": 0.0
        }
    capital = sum(float(c.capital_share or 0.0) for c in project.contributions)
    actual_col = sum(float(c.amount or 0.0) for c in project.collections if c.col_type != "تحويل وارد")
    total_exp = sum(float(e.amount or 0.0) for e in project.expenses)

    col_from_developer = sum(
        float(c.amount or 0.0) for c in project.collections if c.col_type == "دفعة من الشركة المطورة")
    col_from_sales = sum(float(c.amount or 0.0) for c in project.collections if c.col_type in ["مقدم بيع", "قسط بيع"])

    all_dists = db.query(models.Distribution).filter(models.Distribution.source_project_id == project.id).all()
    total_profits_allocated = sum(float(d.amount or 0.0) for d in all_dists if d.op_type in PROFIT_TYPES)
    cash_withdrawn_investors = sum(float(d.amount or 0.0) for d in all_dists if d.op_type in CASH_OUT_TYPES)
    cash_withdrawn_admin = sum(float(d.amount or 0.0) for d in all_dists if d.op_type in ADMIN_CASH_TYPES)

    incoming_transfers = sum(float(d.amount or 0.0) for d in db.query(models.Distribution).filter(
        models.Distribution.target_project_id == project.id,
        models.Distribution.op_type == "تحويل بين المشاريع"
    ).all())
    outgoing_transfers = sum(float(d.amount or 0.0) for d in all_dists if d.op_type == "تحويل بين المشاريع")

    remaining_liquidity = (actual_col + incoming_transfers) - (
                total_exp + outgoing_transfers + cash_withdrawn_investors + cash_withdrawn_admin)
    recovery_pct = round((actual_col / capital * 100), 1) if capital > 0 else 0.0

    return {
        "capital": capital,
        "actual_col": actual_col,
        "col_from_developer": col_from_developer,
        "col_from_sales": col_from_sales,
        "total_exp": total_exp,
        "net_profit": total_profits_allocated,
        "total_profits_allocated": total_profits_allocated,
        "cash_withdrawn_investors": cash_withdrawn_investors,
        "recovery_pct": recovery_pct,
        "remaining_to_recover": max(0.0, capital - actual_col),
        "incoming_transfers": incoming_transfers,
        "outgoing_transfers": outgoing_transfers,
        "remaining_liquidity": remaining_liquidity
    }


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if is_auth(request):
        return RedirectResponse(url="/portal" if get_role(request) == "investor" else "/")
    return templates.TemplateResponse(request=request, name="login.html", context={"error": None})


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    h = hashlib.sha256(password.encode()).hexdigest()
    user = db.query(models.User).filter(models.User.username == username.strip(),
                                        models.User.password_hash == h).first()
    if user:
        request.session["user"] = user.username
        request.session["role"] = user.role or "admin"
        request.session["investor_id"] = user.investor_id
        if user.role == "investor":
            return RedirectResponse(url="/portal", status_code=303)
        return RedirectResponse(url="/", status_code=303)
    return templates.TemplateResponse(request=request, name="login.html", context={"error": "بيانات الدخول غير صحيحة"})


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)


@app.get("/", response_class=HTMLResponse)
def read_dashboard(request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    if get_role(request) == "investor": return RedirectResponse(url="/portal")

    projects = db.query(models.Project).order_by(models.Project.id.desc()).all()
    investors = db.query(models.Investor).order_by(models.Investor.name).all()

    projects_cards = []
    tot_capital = 0.0;
    tot_col = 0.0;
    tot_exp = 0.0;
    tot_profit = 0.0

    for p in projects:
        fin = calc_project_financials(p, db)
        tot_capital += fin["capital"]
        tot_col += fin["actual_col"]
        tot_exp += fin["total_exp"]
        tot_profit += fin["total_profits_allocated"]
        projects_cards.append({"project": p, "fin": fin})

    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "projects_cards": projects_cards,
            "investors": investors,
            "tot_capital": tot_capital,
            "tot_col": tot_col,
            "tot_exp": tot_exp,
            "tot_profit": tot_profit,
            "username": request.session.get("user")
        }
    )


@app.get("/treasury", response_class=HTMLResponse)
def treasury_dashboard(request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    if get_role(request) == "investor": return RedirectResponse(url="/portal")

    projects = db.query(models.Project).order_by(models.Project.id.desc()).all()
    treasury_rows = []
    total_system_liquidity = 0.0
    total_system_collected = 0.0
    total_system_expenses = 0.0
    total_system_withdrawn = 0.0

    for p in projects:
        fin = calc_project_financials(p, db)
        total_system_liquidity += fin["remaining_liquidity"]
        total_system_collected += fin["actual_col"]
        total_system_expenses += fin["total_exp"]
        total_system_withdrawn += fin["cash_withdrawn_investors"]
        treasury_rows.append({"project": p, "fin": fin})

    return templates.TemplateResponse(
        request=request,
        name="treasury.html",
        context={
            "treasury_rows": treasury_rows,
            "total_system_liquidity": total_system_liquidity,
            "total_system_collected": total_system_collected,
            "total_system_expenses": total_system_expenses,
            "total_system_withdrawn": total_system_withdrawn,
            "username": request.session.get("user")
        }
    )


@app.get("/sales", response_class=HTMLResponse)
def sales_dashboard(request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    if get_role(request) == "investor": return RedirectResponse(url="/portal")

    all_sales = db.query(models.Sale).order_by(models.Sale.id.desc()).all()
    all_installments = db.query(models.Installment).order_by(models.Installment.due_date.asc()).all()
    all_projects = db.query(models.Project).order_by(models.Project.name).all()
    available_units = db.query(models.ProjectUnit).filter(models.ProjectUnit.status == "متاحة").all()

    tot_sales_val = sum(float(s.total_price or 0.0) for s in all_sales)
    tot_collected_down = sum(float(s.down_payment or 0.0) for s in all_sales)
    tot_paid_installments = sum(float(i.paid_amount or 0.0) for i in all_installments)
    tot_remaining_installments = sum(
        float((i.amount or 0.0) - (i.paid_amount or 0.0)) for i in all_installments if i.status != "مسدد")
    tot_profit_from_sales = sum(float(s.net_unit_profit or 0.0) for s in all_sales)

    return templates.TemplateResponse(
        request=request,
        name="sales.html",
        context={
            "all_sales": all_sales,
            "all_installments": all_installments,
            "all_projects": all_projects,
            "available_units": available_units,
            "tot_sales_val": tot_sales_val,
            "tot_collected_down": tot_collected_down,
            "tot_paid_installments": tot_paid_installments,
            "tot_remaining_installments": tot_remaining_installments,
            "tot_profit_from_sales": tot_profit_from_sales,
            "username": request.session.get("user")
        }
    )


@app.post("/sales/create")
def create_sale(
        request: Request,
        unit_id: int = Form(0),
        project_id: int = Form(0),
        custom_unit_name: str = Form(""),
        custom_cost_price: float = Form(0.0),
        buyer_name: str = Form(...),
        buyer_phone: str = Form(""),
        total_price: float = Form(...),
        sale_type: str = Form("تقسيط"),
        down_payment: float = Form(0.0),
        installments_count: int = Form(0),
        period_type: str = Form("شهري"),
        first_due_date: str = Form(""),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")

    count = db.query(models.Sale).count() + 1
    sale_code = f"SALE-{count:04d}"

    actual_unit_id = None
    if unit_id > 0:
        unit = db.query(models.ProjectUnit).filter(models.ProjectUnit.id == unit_id).first()
        if unit:
            actual_unit_id = unit.id
            actual_project_id = unit.project_id
            unit_description = f"{unit.unit_name} ({unit.unit_code})"
            cost_price = float(unit.cost_price or 0.0)
            unit.status = "مباعة"
    else:
        actual_project_id = project_id
        unit_description = custom_unit_name.strip()
        cost_price = float(custom_cost_price or 0.0)

    net_unit_profit = max(0.0, float(total_price) - cost_price)

    if sale_type == "كاش":
        down_payment = total_price
        remaining = 0.0
        installments_count = 0
    else:
        remaining = max(0.0, float(total_price) - float(down_payment))

    new_sale = models.Sale(
        sale_code=sale_code,
        project_id=actual_project_id,
        unit_id=actual_unit_id,
        buyer_name=buyer_name.strip(),
        buyer_phone=buyer_phone.strip(),
        unit_description=unit_description,
        sale_date=datetime.utcnow(),
        cost_price=cost_price,
        total_price=float(total_price),
        net_unit_profit=net_unit_profit,
        sale_type=sale_type,
        down_payment=float(down_payment),
        remaining_amount=remaining,
        installments_count=int(installments_count),
        period_type=period_type
    )
    db.add(new_sale)
    db.flush()

    if down_payment > 0:
        c_count = db.query(models.Collection).count() + 1
        col = models.Collection(
            code=f"COL-{c_count:04d}",
            date=datetime.utcnow(),
            project_id=actual_project_id,
            amount=float(down_payment),
            col_type="مقدم بيع",
            description=f"مقدم بيع {unit_description} ({sale_code}) - المشتري: {buyer_name}"
        )
        db.add(col)

    if sale_type == "تقسيط" and installments_count > 0 and remaining > 0:
        inst_amount = round(remaining / int(installments_count), 2)
        start_date = datetime.strptime(first_due_date, "%Y-%m-%d") if first_due_date else datetime.utcnow()
        step = 1
        if period_type == "ربع سنوي":
            step = 3
        elif period_type == "نصف سنوي":
            step = 6
        elif period_type == "سنوي":
            step = 12
        for i in range(int(installments_count)):
            due = add_months(start_date, i * step)
            inst = models.Installment(
                sale_id=new_sale.id,
                project_id=actual_project_id,
                installment_no=i + 1,
                due_date=due,
                amount=inst_amount,
                paid_amount=0.0,
                status="مستحق",
                notes=f"قسط {i + 1} من {installments_count}"
            )
            db.add(inst)

    db.commit()
    return RedirectResponse(url="/sales", status_code=303)


@app.post("/sales/{sale_id}/delete")
def delete_sale(sale_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    sale = db.query(models.Sale).filter(models.Sale.id == sale_id).first()
    if sale:
        if sale.unit:
            sale.unit.status = "متاحة"
        cols = db.query(models.Collection).filter(
            models.Collection.project_id == sale.project_id,
            models.Collection.description.contains(sale.sale_code)
        ).all()
        for c in cols: db.delete(c)
        db.delete(sale)
        db.commit()
    return RedirectResponse(url="/sales", status_code=303)


@app.post("/installments/{installment_id}/pay")
def pay_installment(installment_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    inst = db.query(models.Installment).filter(models.Installment.id == installment_id).first()
    if inst and inst.status != "مسدد":
        inst.status = "مسدد"
        inst.paid_amount = float(inst.amount or 0.0)
        inst.paid_date = datetime.utcnow()
        c_count = db.query(models.Collection).count() + 1
        col = models.Collection(
            code=f"COL-{c_count:04d}",
            date=datetime.utcnow(),
            project_id=inst.project_id,
            amount=inst.paid_amount,
            col_type="قسط بيع",
            description=f"تحصيل قسط رقم {inst.installment_no} ({inst.sale.sale_code} - {inst.sale.unit_description})"
        )
        db.add(col)
        db.commit()
    return RedirectResponse(url="/sales", status_code=303)


@app.post("/installments/{installment_id}/revert")
def revert_installment(installment_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    inst = db.query(models.Installment).filter(models.Installment.id == installment_id).first()
    if inst and inst.status == "مسدد":
        inst.status = "مستحق"
        inst.paid_amount = 0.0
        inst.paid_date = None
        col = db.query(models.Collection).filter(
            models.Collection.project_id == inst.project_id,
            models.Collection.col_type == "قسط بيع",
            models.Collection.description.contains(inst.sale.sale_code)
        ).order_by(models.Collection.id.desc()).first()
        if col: db.delete(col)
        db.commit()
    return RedirectResponse(url="/sales", status_code=303)


@app.get("/projects/{project_id}", response_class=HTMLResponse)
def project_detail(project_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    if get_role(request) == "investor": return RedirectResponse(url="/portal")

    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project: return RedirectResponse(url="/")

    fin = calc_project_financials(project, db)
    project_dists = db.query(models.Distribution).filter(models.Distribution.source_project_id == project.id).all()

    investors_rows = []
    for c in project.contributions:
        cap = float(c.capital_share or 0.0)
        fee_pct = float(c.mgmt_fee_pct or 0.0)
        share_pct = (cap / fin["capital"]) if fin["capital"] > 0 else 0.0
        inv_dists = [d for d in project_dists if d.investor_id == c.investor_id]

        alloc_profit = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in PROFIT_TYPES)
        total_due = cap + alloc_profit
        cash_withdrawn = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in CASH_OUT_TYPES)
        unwithdrawn_profit = max(0.0, alloc_profit - cash_withdrawn)
        remaining = total_due - cash_withdrawn

        investors_rows.append({
            "cont_id": c.id,
            "investor_id": c.investor_id,
            "name": c.investor.name if c.investor else "غير محدد",
            "capital_share": cap,
            "share_pct": round(share_pct * 100, 2),
            "mgmt_fee_pct": round(fee_pct * 100, 1),
            "alloc_profit": alloc_profit,
            "total_due": total_due,
            "cash_withdrawn": cash_withdrawn,
            "unwithdrawn_profit": unwithdrawn_profit,
            "remaining": remaining
        })

    return templates.TemplateResponse(
        request=request,
        name="project_detail.html",
        context={
            "project": project,
            "fin": fin,
            "investors_rows": investors_rows,
            "units": project.units,
            "all_investors": db.query(models.Investor).order_by(models.Investor.name).all(),
            "all_projects": db.query(models.Project).filter(models.Project.id != project.id).all(),
            "collections": project.collections,
            "expenses": project.expenses,
            "distributions": project_dists,
            "username": request.session.get("user")
        }
    )


@app.post("/projects/create")
def create_project(request: Request, code: str = Form(...), name: str = Form(...), project_type: str = Form("عقاري"),
                   status: str = Form("قيد التنفيذ"), total_capital: float = Form(0.0), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    new_p = models.Project(code=code.strip(), name=name.strip(), project_type=project_type.strip(),
                           status=status.strip(), total_capital=float(total_capital or 0.0))
    db.add(new_p)
    db.commit()
    return RedirectResponse(url="/", status_code=303)


@app.post("/projects/{project_id}/update")
def update_project(project_id: int, request: Request, name: str = Form(...), status: str = Form(...),
                   total_capital: float = Form(0.0), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    p = db.query(models.Project).filter(models.Project.id == project_id).first()
    if p:
        p.name = name.strip()
        p.status = status.strip()
        p.total_capital = float(total_capital or 0.0)
        db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/projects/{project_id}/delete")
def delete_project(project_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    p = db.query(models.Project).filter(models.Project.id == project_id).first()
    if p:
        db.delete(p)
        db.commit()
    return RedirectResponse(url="/", status_code=303)


@app.post("/projects/{project_id}/bulk-distribute")
def bulk_distribute(
        project_id: int,
        request: Request,
        bulk_amount: float = Form(...),
        description: str = Form("إقرار دفعة أرباح مستلمة من المطور"),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    project = db.query(models.Project).filter(models.Project.id == project_id).first()
    if not project or bulk_amount <= 0: return RedirectResponse(url=f"/projects/{project_id}")

    total_capital = sum(float(c.capital_share or 0.0) for c in project.contributions)
    actual_col_now = sum(float(c.amount or 0.0) for c in project.collections if c.col_type != "تحويل وارد")
    remaining_capital_to_recover = max(0.0, total_capital - actual_col_now)

    if bulk_amount <= remaining_capital_to_recover:
        pure_profit = 0.0
    else:
        pure_profit = bulk_amount - remaining_capital_to_recover

    ts = int(datetime.utcnow().timestamp())
    for idx, cont in enumerate(project.contributions):
        ratio = (cont.capital_share / total_capital) if total_capital > 0 else 0
        investor_total_share = bulk_amount * ratio
        investor_profit_share = pure_profit * ratio
        mgmt_fee = investor_profit_share * float(cont.mgmt_fee_pct or 0.0) if cont.mgmt_fee_pct else 0.0
        net_due = investor_total_share - mgmt_fee

        note = f"{description.strip()}"
        if mgmt_fee > 0:
            note += f" (شامل خصم عمولة إدارة {mgmt_fee:,.0f} ج.م على الأرباح الزائدة عن رأس المال)"

        dist = models.Distribution(
            code=f"DST-{ts}-{idx + 1}",
            date=datetime.utcnow(),
            op_type="أرباح معتمدة",
            source_project_id=project_id,
            investor_id=cont.investor_id,
            amount=net_due,
            description=note
        )
        db.add(dist)

    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/projects/{project_id}/withdrawals/add")
def add_withdrawal(
        project_id: int,
        request: Request,
        investor_id: int = Form(...),
        amount: float = Form(...),
        description: str = Form("صرف مسحوبات أرباح نقداً"),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    count = db.query(models.Distribution).count() + 1
    dist = models.Distribution(
        code=f"DST-{count:04d}",
        date=datetime.utcnow(),
        op_type="مسحوب للمساهم",
        source_project_id=project_id,
        investor_id=investor_id,
        amount=float(amount),
        description=description.strip()
    )
    db.add(dist)
    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/transfer-between-projects")
def transfer_between_projects(
        request: Request,
        source_project_id: int = Form(...),
        target_project_id: int = Form(...),
        investor_id: int = Form(...),
        amount: float = Form(...),
        description: str = Form("تحويل مستحقات لمشروع آخر"),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    dist = models.Distribution(
        code=f"TRF-{int(datetime.utcnow().timestamp())}",
        date=datetime.utcnow(),
        op_type="تحويل بين المشاريع",
        source_project_id=source_project_id,
        target_project_id=target_project_id,
        investor_id=investor_id,
        amount=float(amount),
        description=description.strip()
    )
    db.add(dist)
    target_cont = db.query(models.Contribution).filter(models.Contribution.project_id == target_project_id,
                                                       models.Contribution.investor_id == investor_id).first()
    if target_cont:
        target_cont.capital_share += float(amount)
    else:
        db.add(models.Contribution(project_id=target_project_id, investor_id=investor_id, capital_share=float(amount),
                                   mgmt_fee_pct=0.0))
    db.commit()
    return RedirectResponse(url=f"/projects/{source_project_id}", status_code=303)


@app.post("/distributions/{dist_id}/delete")
def delete_distribution(dist_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    dist = db.query(models.Distribution).filter(models.Distribution.id == dist_id).first()
    if dist:
        p_id = dist.source_project_id
        db.delete(dist)
        db.commit()
        return RedirectResponse(url=f"/projects/{p_id}", status_code=303)
    return RedirectResponse(url="/")


@app.post("/projects/{project_id}/units/add")
def add_project_unit(
        project_id: int,
        request: Request,
        unit_code: str = Form(...),
        unit_name: str = Form(...),
        unit_type: str = Form("شقة"),
        area: float = Form(0.0),
        cost_price: float = Form(0.0),
        target_price: float = Form(0.0),
        notes: str = Form(""),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    unit = models.ProjectUnit(
        project_id=project_id,
        unit_code=unit_code.strip(),
        unit_name=unit_name.strip(),
        unit_type=unit_type.strip(),
        area=float(area or 0.0),
        cost_price=float(cost_price or 0.0),
        target_price=float(target_price or 0.0),
        status="متاحة",
        notes=notes.strip()
    )
    db.add(unit)
    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/units/{unit_id}/delete")
def delete_project_unit(unit_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    unit = db.query(models.ProjectUnit).filter(models.ProjectUnit.id == unit_id).first()
    if unit:
        p_id = unit.project_id
        db.delete(unit)
        db.commit()
        return RedirectResponse(url=f"/projects/{p_id}", status_code=303)
    return RedirectResponse(url="/")


@app.post("/contributions/{cont_id}/edit")
def edit_contribution(cont_id: int, request: Request, capital_share: float = Form(...), mgmt_fee_pct: str = Form("0"),
                      redirect_to: str = Form(""), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    cont = db.query(models.Contribution).filter(models.Contribution.id == cont_id).first()
    if cont:
        try:
            fee_num = float(mgmt_fee_pct) if str(mgmt_fee_pct).strip() else 0.0
        except Exception:
            fee_num = 0.0
        cont.capital_share = float(capital_share)
        cont.mgmt_fee_pct = max(0.0, fee_num) / 100.0
        db.commit()
        if redirect_to: return RedirectResponse(url=redirect_to, status_code=303)
        return RedirectResponse(url=f"/projects/{cont.project_id}", status_code=303)
    return RedirectResponse(url="/")


@app.post("/contributions/{cont_id}/delete")
def delete_contribution(cont_id: int, request: Request, redirect_to: str = Form(""), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    cont = db.query(models.Contribution).filter(models.Contribution.id == cont_id).first()
    if cont:
        p_id = cont.project_id
        db.delete(cont)
        db.commit()
        if redirect_to: return RedirectResponse(url=redirect_to, status_code=303)
        return RedirectResponse(url=f"/projects/{p_id}", status_code=303)
    return RedirectResponse(url="/")


@app.post("/projects/{project_id}/collections/add")
def add_collection(project_id: int, request: Request, amount: float = Form(...), col_type: str = Form("دفعة إضافية"),
                   description: str = Form(""), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    count = db.query(models.Collection).count() + 1
    col = models.Collection(code=f"COL-{count:04d}", date=datetime.utcnow(), project_id=project_id,
                            amount=float(amount), col_type=col_type.strip(), description=description.strip())
    db.add(col)
    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/projects/{project_id}/expenses/add")
def add_expense(project_id: int, request: Request, amount: float = Form(...), exp_type: str = Form("مصاريف تشغيل"),
                description: str = Form(""), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    exp = models.Expense(date=datetime.utcnow(), project_id=project_id, exp_type=exp_type.strip(),
                         description=description.strip(), amount=float(amount))
    db.add(exp)
    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


@app.post("/projects/{project_id}/contributions/add")
def add_contribution(project_id: int, request: Request, investor_id: int = Form(...), capital_share: float = Form(...),
                     mgmt_fee_pct: str = Form("0"), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    try:
        fee_num = float(mgmt_fee_pct) if str(mgmt_fee_pct).strip() else 0.0
    except Exception:
        fee_num = 0.0
    fee_pct = max(0.0, fee_num) / 100.0
    existing = db.query(models.Contribution).filter(models.Contribution.project_id == project_id,
                                                    models.Contribution.investor_id == investor_id).first()
    if existing:
        existing.capital_share += float(capital_share)
        existing.mgmt_fee_pct = fee_pct
    else:
        new_c = models.Contribution(project_id=project_id, investor_id=investor_id, capital_share=float(capital_share),
                                    mgmt_fee_pct=fee_pct)
        db.add(new_c)
    db.commit()
    return RedirectResponse(url=f"/projects/{project_id}", status_code=303)


# ----------------- كشوف حسابات المساهمين -----------------
@app.get("/investors", response_class=HTMLResponse)
def investors_directory(request: Request, investor_id: int = None, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    if get_role(request) == "investor": return RedirectResponse(url="/portal")

    investors = db.query(models.Investor).order_by(models.Investor.name).all()
    all_projects = db.query(models.Project).order_by(models.Project.name).all()

    selected_inv = None
    portfolio_rows = []
    transactions_history = []
    tot_cap = 0.0;
    tot_profit = 0.0;
    tot_due = 0.0;
    tot_withdrawn = 0.0;
    tot_remaining = 0.0
    whatsapp_url = None
    existing_user = None

    if investor_id:
        selected_inv = db.query(models.Investor).filter(models.Investor.id == investor_id).first()
        if selected_inv:
            try:
                existing_user = db.query(models.User).filter(models.User.investor_id == selected_inv.id).first()
            except Exception:
                existing_user = None

            wa_projects_details = []

            for cont in getattr(selected_inv, "contributions", []):
                if not getattr(cont, "project", None):
                    continue
                p = cont.project
                p_fin = calc_project_financials(p, db)
                cap = float(cont.capital_share or 0.0)
                fee_pct = float(cont.mgmt_fee_pct or 0.0)
                share_pct = (cap / p_fin["capital"]) if p_fin["capital"] > 0 else 0.0

                inv_dists = db.query(models.Distribution).filter(
                    models.Distribution.source_project_id == p.id,
                    models.Distribution.investor_id == selected_inv.id
                ).all()

                alloc_profit = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in PROFIT_TYPES)
                total_due = cap + alloc_profit
                cash_withdrawn = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in CASH_OUT_TYPES)
                remaining = total_due - cash_withdrawn
                unwithdrawn_profit = max(0.0, alloc_profit - cash_withdrawn)

                tot_cap += cap
                tot_profit += alloc_profit
                tot_due += total_due
                tot_withdrawn += cash_withdrawn
                tot_remaining += remaining

                portfolio_rows.append({
                    "cont_id": cont.id,
                    "project_id": p.id,
                    "project_name": p.name,
                    "capital": cap,
                    "share_pct": round(share_pct * 100, 2),
                    "mgmt_fee_pct": round(fee_pct * 100, 1),
                    "alloc_profit": alloc_profit,
                    "total_due": total_due,
                    "cash_withdrawn": cash_withdrawn,
                    "unwithdrawn_profit": unwithdrawn_profit,
                    "remaining": remaining
                })

                wa_projects_details.append(
                    f"🔹 *مشروع: {p.name}*\n"
                    f"  ▫️ رأس المال: {cap:,.0f} ج.م ({round(share_pct * 100, 2)}%)\n"
                    f"  ▫️ صافي الأرباح المقررة: {alloc_profit:,.0f} ج.م\n"
                    f"  ▫️ إجمالي المستحق: {total_due:,.0f} ج.م\n"
                    f"  ▫️ المسحوبات المستلمة: {cash_withdrawn:,.0f} ج.م\n"
                    f"  ▫️ الرصيد المتبقي المستحق: *{remaining:,.0f} ج.م*"
                )

                transactions_history.append({
                    "date": "بداية المساهمة",
                    "project": p.name,
                    "type": "حصة رأس مال",
                    "inflow": cap,
                    "outflow": 0.0,
                    "notes": f"نسبة مساهمة {round(share_pct * 100, 2)}%"
                })

            all_dists = db.query(models.Distribution).filter(
                models.Distribution.investor_id == selected_inv.id).order_by(models.Distribution.date.desc()).all()
            for d in all_dists:
                source_name = d.source_project.name if getattr(d, "source_project", None) else "عام"
                is_cash_out = d.op_type in CASH_OUT_TYPES
                amt = float(d.amount or 0.0)
                transactions_history.append({
                    "date": d.date.strftime("%Y-%m-%d") if getattr(d, "date", None) else "-",
                    "project": source_name,
                    "type": d.op_type or "",
                    "inflow": amt if d.op_type in PROFIT_TYPES else 0.0,
                    "outflow": amt if is_cash_out else 0.0,
                    "notes": d.description or ""
                })

            phone_clean = ''.join(filter(str.isdigit, selected_inv.phone or ""))
            if phone_clean.startswith("01"): phone_clean = "20" + phone_clean[1:]

            projects_summary_text = "\n\n".join(wa_projects_details)
            wa_text = (
                f"السلام عليكم ورحمة الله وبركاته،\n"
                f"أهلاً بحضرتك أستاذ/ *{selected_inv.name}* 🤝\n\n"
                f"📊 *كشف حساب استثماري رسمي حتى تاريخ اليوم:*\n"
                f"━━━━━━━━━━━━━━━━━━━\n\n"
                f"{projects_summary_text}\n\n"
                f"━━━━━━━━━━━━━━━━━━━\n"
                f"📈 *الإجماليات الكلية لكافة المشاريع:*\n"
                f"▫️ إجمالي رأس المال: {tot_cap:,.0f} ج.م\n"
                f"▫️ إجمالي الأرباح المقررة: {tot_profit:,.0f} ج.م\n"
                f"▫️ إجمالي المسحوبات المستلمة: {tot_withdrawn:,.0f} ج.م\n"
                f"▫️ *إجمالي الرصيد المتبقي المستحق: {tot_remaining:,.0f} ج.م*\n"
                f"━━━━━━━━━━━━━━━━━━━\n\n"
                f"مع خالص التمنيات بالتوفيق والازدهار الدائم."
            )
            whatsapp_url = f"https://wa.me/{phone_clean}?text={urllib.parse.quote(wa_text)}" if phone_clean else f"https://wa.me/?text={urllib.parse.quote(wa_text)}"

    return templates.TemplateResponse(
        request=request,
        name="investors.html",
        context={
            "investors": investors,
            "all_projects": all_projects,
            "selected_inv": selected_inv,
            "portfolio_rows": portfolio_rows,
            "transactions_history": transactions_history,
            "tot_cap": tot_cap,
            "tot_profit": tot_profit,
            "tot_due": tot_due,
            "tot_withdrawn": tot_withdrawn,
            "tot_remaining": tot_remaining,
            "whatsapp_url": whatsapp_url,
            "existing_user": existing_user,
            "username": request.session.get("user")
        }
    )


@app.post("/investors/create")
def create_investor(request: Request, name: str = Form(...), phone: str = Form(""), notes: str = Form(""),
                    db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    count = db.query(models.Investor).count() + 1
    new_inv = models.Investor(code=f"INV-{count:03d}", name=name.strip(), phone=phone.strip(), notes=notes.strip())
    db.add(new_inv)
    db.commit()
    return RedirectResponse(url=f"/investors?investor_id={new_inv.id}", status_code=303)


@app.post("/investors/{investor_id}/update")
def update_investor(investor_id: int, request: Request, name: str = Form(...), phone: str = Form(""),
                    notes: str = Form(""), db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    inv = db.query(models.Investor).filter(models.Investor.id == investor_id).first()
    if inv:
        inv.name = name.strip();
        inv.phone = phone.strip();
        inv.notes = notes.strip();
        db.commit()
    return RedirectResponse(url=f"/investors?investor_id={investor_id}", status_code=303)


# إنشاء أو تعديل بيانات دخول المساهم
@app.post("/investors/{investor_id}/create-account")
def create_or_update_investor_account(
        investor_id: int,
        request: Request,
        username: str = Form(...),
        password: str = Form(""),
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    user = db.query(models.User).filter(models.User.investor_id == investor_id).first()
    if user:
        user.username = username.strip()
        if password.strip():
            user.password_hash = hashlib.sha256(password.strip().encode()).hexdigest()
    else:
        pwd = password.strip() if password.strip() else "123456"
        h = hashlib.sha256(pwd.encode()).hexdigest()
        new_u = models.User(username=username.strip(), password_hash=h, role="investor", investor_id=investor_id)
        db.add(new_u)
    db.commit()
    return RedirectResponse(url=f"/investors?investor_id={investor_id}", status_code=303)


# إلغاء / حذف حساب الدخول للمساهم
@app.post("/investors/{investor_id}/delete-account")
def delete_investor_account(
        investor_id: int,
        request: Request,
        db: Session = Depends(get_db)
):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    user = db.query(models.User).filter(models.User.investor_id == investor_id).first()
    if user:
        db.delete(user)
        db.commit()
    return RedirectResponse(url=f"/investors?investor_id={investor_id}", status_code=303)


# تصدير كشف حساب المساهم Excel (معالجة الأحرف والرموز الخاصة بأمان 100%)
@app.get("/investors/{investor_id}/export-excel")
def export_investor_excel(investor_id: int, request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    inv = db.query(models.Investor).filter(models.Investor.id == investor_id).first()
    if not inv: return RedirectResponse(url="/investors")

    wb = openpyxl.Workbook()
    ws = wb.active

    # تنظيف اسم الورقة لمنع خطأ الرموز الممنوعة
    safe_sheet_title = f"كشف حساب - {inv.name}"
    for ch in ['\\', '/', '?', '*', ':', '[', ']']:
        safe_sheet_title = safe_sheet_title.replace(ch, '-')
    ws.title = safe_sheet_title[:30]
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    ws["A1"] = f"كشف حساب استثماري للمساهم: {inv.name}"
    ws["A1"].font = Font(name="Segoe UI", size=14, bold=True)
    ws["A2"] = f"تاريخ الاستخراج: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Segoe UI", size=10, italic=True)

    ws["A4"] = "أولاً: ملخص المساهمات والأرباح والمسحوبات"
    ws["A4"].font = Font(name="Segoe UI", size=12, bold=True)

    headers_p = ["اسم المشروع", "رأس المال الأصلي", "نسبة الملكية", "نسبة الإدارة", "صافي الأرباح المقررة",
                 "إجمالي المستحق", "المسحوبات المستلمة", "الرصيد المتبقي"]
    for idx, h in enumerate(headers_p, 1):
        cell = ws.cell(row=5, column=idx, value=h)
        cell.fill = header_fill;
        cell.font = header_font;
        cell.alignment = Alignment(horizontal="center")

    r = 6
    for cont in getattr(inv, "contributions", []):
        if not getattr(cont, "project", None): continue
        p_fin = calc_project_financials(cont.project, db)
        cap = float(cont.capital_share or 0.0)
        fee_pct = float(cont.mgmt_fee_pct or 0.0)
        share_pct = (cap / p_fin["capital"]) if p_fin["capital"] > 0 else 0.0
        inv_dists = db.query(models.Distribution).filter(
            models.Distribution.source_project_id == cont.project.id,
            models.Distribution.investor_id == inv.id
        ).all()
        alloc_profit = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in PROFIT_TYPES)
        total_due = cap + alloc_profit
        cash_withdrawn = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in CASH_OUT_TYPES)
        rem = total_due - cash_withdrawn

        ws.cell(row=r, column=1, value=cont.project.name)
        ws.cell(row=r, column=2, value=cap)
        ws.cell(row=r, column=3, value=f"{round(share_pct * 100, 2)}%")
        ws.cell(row=r, column=4, value=f"{round(fee_pct * 100, 1)}%")
        ws.cell(row=r, column=5, value=alloc_profit)
        ws.cell(row=r, column=6, value=total_due)
        ws.cell(row=r, column=7, value=cash_withdrawn)
        ws.cell(row=r, column=8, value=rem)
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="ثانياً: سجل الحركات والمسحوبات والتسويات").font = Font(name="Segoe UI", size=12,
                                                                                           bold=True)
    r += 1
    headers_tx = ["التاريخ", "المشروع المصدر", "نوع الحركة", "المبلغ", "البيان والتفاصيل"]
    for idx, h in enumerate(headers_tx, 1):
        cell = ws.cell(row=r, column=idx, value=h)
        cell.fill = header_fill;
        cell.font = header_font;
        cell.alignment = Alignment(horizontal="center")

    r += 1
    dists = db.query(models.Distribution).filter(models.Distribution.investor_id == inv.id).order_by(
        models.Distribution.date.desc()).all()
    for d in dists:
        ws.cell(row=r, column=1, value=d.date.strftime("%Y-%m-%d") if getattr(d, "date", None) else "")
        ws.cell(row=r, column=2, value=d.source_project.name if getattr(d, "source_project", None) else "")
        ws.cell(row=r, column=3, value=d.op_type or "")
        ws.cell(row=r, column=4, value=float(d.amount or 0.0))
        ws.cell(row=r, column=5, value=d.description or "")
        r += 1

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    # تشفير الترويسة القياسي الداعم للعربية
    safe_ascii_filename = f"Statement_INV_{inv.id}.xlsx"
    clean_inv_name = inv.name.replace("/", "-").replace("\\", "-")
    encoded_utf8_filename = urllib.parse.quote(f"كشف_حساب_{clean_inv_name}.xlsx")
    content_disposition = f'attachment; filename="{safe_ascii_filename}"; filename*=UTF-8\'\'{encoded_utf8_filename}'

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition}
    )


# ----------------- بوابة المساهم المعزولة -----------------
@app.get("/portal", response_class=HTMLResponse)
def investor_portal(request: Request, db: Session = Depends(get_db)):
    if not is_auth(request): return RedirectResponse(url="/login")
    inv_id = request.session.get("investor_id")
    current_username = request.session.get("user")

    inv = None
    if inv_id:
        inv = db.query(models.Investor).filter(models.Investor.id == inv_id).first()
    if not inv:
        u = db.query(models.User).filter(models.User.username == current_username).first()
        if u and u.investor_id:
            inv = db.query(models.Investor).filter(models.Investor.id == u.investor_id).first()

    if not inv:
        if get_role(request) == "admin": return RedirectResponse(url="/")
        return RedirectResponse(url="/login")

    portfolio_rows = []
    transactions_history = []
    tot_cap = 0.0;
    tot_profit = 0.0;
    tot_due = 0.0;
    tot_withdrawn = 0.0;
    tot_remaining = 0.0

    for cont in getattr(inv, "contributions", []):
        if not getattr(cont, "project", None): continue
        p_fin = calc_project_financials(cont.project, db)
        cap = float(cont.capital_share or 0.0)
        fee_pct = float(cont.mgmt_fee_pct or 0.0)
        share_pct = (cap / p_fin["capital"]) if p_fin["capital"] > 0 else 0.0
        inv_dists = db.query(models.Distribution).filter(models.Distribution.source_project_id == cont.project.id,
                                                         models.Distribution.investor_id == inv.id).all()
        alloc_profit = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in PROFIT_TYPES)
        total_due = cap + alloc_profit
        cash_withdrawn = sum(float(d.amount or 0.0) for d in inv_dists if d.op_type in CASH_OUT_TYPES)
        remaining = total_due - cash_withdrawn
        unwithdrawn_profit = max(0.0, alloc_profit - cash_withdrawn)

        tot_cap += cap;
        tot_profit += alloc_profit;
        tot_due += total_due;
        tot_withdrawn += cash_withdrawn;
        tot_remaining += remaining
        portfolio_rows.append({
            "project_name": cont.project.name,
            "project_status": cont.project.status,
            "capital": cap,
            "share_pct": round(share_pct * 100, 2),
            "alloc_profit": alloc_profit,
            "net_profit": alloc_profit,
            "total_due": total_due,
            "cash_withdrawn": cash_withdrawn,
            "unwithdrawn_profit": unwithdrawn_profit,
            "remaining": remaining
        })
        transactions_history.append({
            "date": "بداية المساهمة", "project": cont.project.name, "type": "حصة مساهمة",
            "inflow": cap, "outflow": 0.0, "notes": f"نسبة الملكية {round(share_pct * 100, 2)}%"
        })

    all_dists = db.query(models.Distribution).filter(models.Distribution.investor_id == inv.id).order_by(
        models.Distribution.date.desc()).all()
    for d in all_dists:
        source_name = d.source_project.name if getattr(d, "source_project", None) else "عام"
        is_cash_out = d.op_type in CASH_OUT_TYPES
        amt = float(d.amount or 0.0)
        transactions_history.append({
            "date": d.date.strftime("%Y-%m-%d") if getattr(d, "date", None) else "-", "project": source_name,
            "type": d.op_type or "", "inflow": amt if d.op_type in PROFIT_TYPES else 0.0,
            "outflow": amt if is_cash_out else 0.0, "notes": d.description or ""
        })

    return templates.TemplateResponse(
        request=request, name="investor_portal.html",
        context={
            "inv": inv, "portfolio_rows": portfolio_rows, "transactions_history": transactions_history,
            "tot_cap": tot_cap, "tot_profit": tot_profit, "tot_due": tot_due,
            "tot_withdrawn": tot_withdrawn, "tot_remaining": tot_remaining, "username": inv.name
        }
    )


@app.post("/system/reset-all-data")
def reset_all_data(request: Request, db: Session = Depends(get_db)):
    if not is_auth(request) or get_role(request) != "admin": return RedirectResponse(url="/login")
    db.query(models.Installment).delete()
    db.query(models.Sale).delete()
    db.query(models.ProjectUnit).delete()
    db.query(models.Distribution).delete()
    db.query(models.Expense).delete()
    db.query(models.Collection).delete()
    db.query(models.Contribution).delete()
    db.query(models.Project).delete()
    db.query(models.Investor).delete()
    db.query(models.User).filter(models.User.role != "admin").delete()
    db.commit()
    return RedirectResponse(url="/", status_code=303)