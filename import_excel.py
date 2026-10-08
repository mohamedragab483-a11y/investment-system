import openpyxl
import datetime
import os
import hashlib
from database import engine, SessionLocal, Base
import models

# إعادة تهيئة الجداول
Base.metadata.drop_all(bind=engine)
Base.metadata.create_all(bind=engine)

db = SessionLocal()

# إنشاء مستخدم الإدارة
admin_hash = hashlib.sha256("admin123".encode()).hexdigest()
db.add(models.User(username="admin", password_hash=admin_hash))
db.commit()

# البحث عن ملف الإكسيل
excel_file = "data.xlsx"
if not os.path.exists(excel_file):
    excel_file = "data.xlsx.xlsx"

wb = openpyxl.load_workbook(excel_file, data_only=True)
print("جاري قراءة ملف الإكسيل واستيراد البيانات...")

# 1. استيراد المشاريع
ws_proj = wb['المشاريع']
projects_map = {}
for r in range(3, ws_proj.max_row + 1):
    code = ws_proj.cell(r, 1).value
    name = ws_proj.cell(r, 2).value
    if not name: continue
    proj_type = ws_proj.cell(r, 3).value or "عقاري"
    status = ws_proj.cell(r, 6).value or "قيد التنفيذ"
    capital = ws_proj.cell(r, 7).value or 0.0

    p = models.Project(
        code=str(code) if code else f"PRJ-{r-2:03d}",
        name=str(name).strip(),
        project_type=str(proj_type).strip(),
        status=str(status).strip(),
        total_capital=float(capital) if isinstance(capital, (int, float)) else 0.0
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    projects_map[p.name] = p

print(f"✓ تم استيراد {len(projects_map)} مشاريع.")

# 2. استيراد المساهمين
ws_inv = wb['المساهمون']
investors_map = {}
for r in range(3, ws_inv.max_row + 1):
    code = ws_inv.cell(r, 1).value
    name = ws_inv.cell(r, 2).value
    if not name: continue
    phone = ws_inv.cell(r, 3).value
    notes = ws_inv.cell(r, 4).value

    inv = models.Investor(
        code=str(code) if code else f"INV-{r-2:03d}",
        name=str(name).strip(),
        phone=str(phone or ""),
        notes=str(notes or "")
    )
    db.add(inv)
    db.commit()
    db.refresh(inv)
    investors_map[inv.name] = inv

print(f"✓ تم استيراد {len(investors_map)} مساهم.")

# 3. استيراد مساهمات المشاريع
ws_cont = wb['مساهمات المشاريع']
cont_count = 0
for r in range(3, ws_cont.max_row + 1):
    proj_name = ws_cont.cell(r, 2).value
    inv_name = ws_cont.cell(r, 4).value
    amount = ws_cont.cell(r, 5).value
    mgmt_fee = ws_cont.cell(r, 7).value
    if not proj_name or not inv_name: continue

    p = projects_map.get(str(proj_name).strip())
    inv = investors_map.get(str(inv_name).strip())
    if not inv:
        inv = models.Investor(name=str(inv_name).strip())
        db.add(inv)
        db.commit()
        db.refresh(inv)
        investors_map[inv.name] = inv

    if p and inv:
        cont = models.Contribution(
            project_id=p.id,
            investor_id=inv.id,
            capital_share=float(amount) if isinstance(amount, (int, float)) else 0.0,
            mgmt_fee_pct=float(mgmt_fee) if isinstance(mgmt_fee, (int, float)) else 0.20
        )
        db.add(cont)
        cont_count += 1

db.commit()
print(f"✓ تم استيراد {cont_count} مساهمة وحصص ملكية.")

# 4. استيراد التحصيلات
ws_col = wb['التحصيلات']
col_count = 0
for r in range(3, ws_col.max_row + 1):
    op_num = ws_col.cell(r, 1).value
    dt = ws_col.cell(r, 2).value
    proj_name = ws_col.cell(r, 3).value
    amount = ws_col.cell(r, 5).value
    col_type = ws_col.cell(r, 6).value
    desc = ws_col.cell(r, 7).value
    if not proj_name or not amount: continue

    p = projects_map.get(str(proj_name).strip())
    if p:
        col = models.Collection(
            code=str(op_num or ""),
            date=dt if isinstance(dt, datetime.datetime) else datetime.datetime.utcnow(),
            project_id=p.id,
            amount=float(amount),
            col_type=str(col_type or "تحصيل"),
            description=str(desc or "")
        )
        db.add(col)
        col_count += 1

db.commit()
print(f"✓ تم استيراد {col_count} عملية تحصيل.")

# 5. استيراد المصروفات
ws_exp = wb['المصروفات']
exp_count = 0
for r in range(3, ws_exp.max_row + 1):
    dt = ws_exp.cell(r, 1).value
    proj_name = ws_exp.cell(r, 2).value
    exp_type = ws_exp.cell(r, 3).value
    desc = ws_exp.cell(r, 4).value
    amount = ws_exp.cell(r, 5).value
    if not proj_name or not amount: continue

    p = None
    for k, v in projects_map.items():
        if str(proj_name).strip() in k or k in str(proj_name).strip():
            p = v
            break
    if p:
        exp = models.Expense(
            date=dt if isinstance(dt, datetime.datetime) else datetime.datetime.utcnow(),
            project_id=p.id,
            exp_type=str(exp_type or "مصاريف تشغيل"),
            description=str(desc or ""),
            amount=float(amount)
        )
        db.add(exp)
        exp_count += 1

db.commit()
print(f"✓ تم استيراد {exp_count} عملية صرف.")

# 6. استيراد التوزيعات والتسويات
ws_dist = wb['التوزيعات والتسويات']
dist_count = 0
for r in range(3, ws_dist.max_row + 1):
    op_num = ws_dist.cell(r, 1).value
    dt = ws_dist.cell(r, 2).value
    op_type = ws_dist.cell(r, 3).value
    from_proj = ws_dist.cell(r, 4).value
    to_proj = ws_dist.cell(r, 5).value
    inv_name = ws_dist.cell(r, 6).value
    amount = ws_dist.cell(r, 7).value
    desc = ws_dist.cell(r, 10).value
    if not op_type or not amount: continue

    src_p = projects_map.get(str(from_proj).strip()) if from_proj else None
    tgt_p = projects_map.get(str(to_proj).strip()) if to_proj else None
    inv = investors_map.get(str(inv_name).strip()) if inv_name else None

    dist = models.Distribution(
        code=str(op_num or ""),
        date=dt if isinstance(dt, datetime.datetime) else datetime.datetime.utcnow(),
        op_type=str(op_type).strip(),
        source_project_id=src_p.id if src_p else None,
        target_project_id=tgt_p.id if tgt_p else None,
        investor_id=inv.id if inv else None,
        amount=float(amount),
        description=str(desc or "")
    )
    db.add(dist)
    dist_count += 1

db.commit()
print(f"✓ تم استيراد {dist_count} عملية توزيع ومسحوبات.")

db.close()
print("\n اكتمل استيراد ملف data.xlsx بالكامل وبنجاح تام في قاعدة البيانات!")